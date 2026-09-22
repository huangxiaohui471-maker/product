# -*- coding: utf-8 -*-
"""生成《全球选品平台_字段口径表_v1.5.xlsx》

v1.5 的改动（对应「国家市场 + 品类细分 + 成分库 + 机会洞察补充」这一轮升级）：
  - 字段 35 → 39：商品主数据新增「二级类目 / 三级类目 / 类目ID」，平台数据新增「国家/地区」
  - 新增 Sheet「品类体系」：平台三级类目树 + 本库功能子类 + 三级类目分布
  - 新增 Sheet「成分库」：43 个成分 × 命中商品数 × 代表类目
  - 洞察规则 5 条 → 9 条：新增「市场机会 / 增速品类 / 新品机会 / 成分机会」
  - 页面看板新增「国家/地区市场」对比表与「二级类目增速榜」；品类库视图改为双 Tab（品类细分树 / 成分库）

v1.4 的改动（对应口碑数据落库）：
  - 新增第二张云表「自有店铺口碑库」（21 字段，37 个本店在售商品）
  - 页面新增第 5 个视图「口碑诊断」，读第二张表
  - 「评分 / 评价数」由 FastMoss 详情页自动抓到（215/250 条有评分）
  - 360 条真实记录补齐 AI 增强：外文标题译中文品名、选品笔记追加「AI 速读」
  - 评价通道现状：抖音后台（罗盘·用户原声）已走通；天猫 / 小红书 / TikTok 未登录，抓不到

v1.3 的改动（对应机会洞察规则改版）：
  - 窗口期：去掉「关联达人数 ≤ 15」「销售额 ≥ 100 万」两个门槛，只看环比增速
  - 改良机会：由「退货率 ≥ 10% 且 销量 ≥ 3 万」改为「差评关键词非空」
  - 风险提醒：由「退货率 ≥ 12% 或 环比 < 0」改为「环比 < 0 或 差评关键词含高危词」

v1.2 的改动：
  - 新增「数据源接入台账 / 市场可走通性 / 洞察规则与前提 / 字段可得性」四张表
  - 字段清单改为从线上页面自动抽取，避免文档与产品漂移

重要：本机 recalc 引擎不可用（无 LibreOffice、无 python formulas 包），
      全表不使用任何 Excel 公式。
"""
import collections
import json
import os
import re

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE = "/Users/huizidemac/WorkBuddy/2026-09-19-21-23-47"
BUILD = BASE + "/_build"
OUT = BASE + "/全球选品平台_字段口径表_v1.5.xlsx"
HTML = BASE + "/全球选品平台.html"

HEAD_FILL = PatternFill("solid", fgColor="1F3A5F")
HEAD_FONT = Font(name="PingFang SC", size=11, bold=True, color="FFFFFF")
CELL_FONT = Font(name="PingFang SC", size=10.5, color="1D1D1F")
GRP_FONT = Font(name="PingFang SC", size=10.5, bold=True, color="0B4F9E")
GRP_FILL = PatternFill("solid", fgColor="EAF2FC")
WARN_FONT = Font(name="PingFang SC", size=10.5, color="A52214")
OK_FONT = Font(name="PingFang SC", size=10.5, color="14663A")
THIN = Side(style="thin", color="D8DCE2")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
TOP_WRAP = Alignment(vertical="top", wrap_text=True)


def style_sheet(ws, widths, n_cols, freeze="A2"):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for c in range(1, n_cols + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = HEAD_FILL
        cell.font = HEAD_FONT
        cell.alignment = Alignment(vertical="center", horizontal="left", wrap_text=True)
        cell.border = BORDER
    ws.row_dimensions[1].height = 28
    ws.freeze_panes = freeze


def fill_rows(ws, rows, start=2, group_col=None, colorize=None, heights=None):
    r = start
    for row in rows:
        is_group = group_col is not None and isinstance(row[0], str) and row[0].startswith("###")
        for c, val in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c, value=(val[3:] if is_group else val))
            cell.font = GRP_FONT if is_group else CELL_FONT
            cell.alignment = TOP_WRAP
            cell.border = BORDER
            if is_group:
                cell.fill = GRP_FILL
        if colorize and not is_group:
            colorize(ws, r, row)
        if heights:
            ws.row_dimensions[r].height = heights
        r += 1
    return r


# ---------- 从页面抽取 FIELDS，保证文档与产品同源 ----------
def load_fields():
    src = open(HTML, encoding="utf-8").read()
    i = src.find("var FIELDS = [")
    j = src.find("\n];", i)
    body = src[i:j]
    out = []
    for obj in re.findall(r"\{[^{}]*\}", body):
        def gv(key):
            m = re.search(r"\b%s:\s*'([^']*)'" % key, obj)
            return m.group(1) if m else None
        n, t, grp, s = gv("n"), gv("t"), gv("g"), gv("src")
        if not n or not grp:
            continue
        om = re.search(r"opts:\s*\[([^\]]*)\]", obj)
        opts = re.findall(r"'([^']+)'", om.group(1)) if om else []
        out.append({"n": n, "t": t, "g": grp, "src": s or "auto", "opts": opts})
    return out


FIELDS = load_fields()

SRC_LABEL = {"auto": "平台可导", "half": "部分可导", "manual": "需人工标"}
SRC_OWNER = {
    "商品名称": "三平台均有", "品牌": "三平台均有", "商品ID": "三平台均有",
    "品类": "由三级类目归并成 6 个业务大类（护肤/彩妆/个护/身体/香氛/工具）",
    "细分品类": "平台二级 / 三级类目（旧字段，保留兼容）",
    "二级类目": "由三级类目归并成的功能子类（底妆 / 面部清洁 / 身体清洁…），46 种",
    "三级类目": "FastMoss 类目树最细一级（沐浴露与香皂 / 遮瑕与粉底…），71 种",
    "类目ID": "FastMoss filterInfo 的 cid 链（如 14/848648/601554），可按它精确复现取数",
    "国家/地区": "FastMoss region + 罗盘 / 蝉妈妈固定中国；比「所属市场」大区细一级",
    "价格": "三平台均有", "上市日期": "三平台均有",
    "数据来源": "导入时按所选源自动打标",
    "所属市场": "导入时按所选源自动带出（中国 / 韩国 / 日本 / 欧美 / 东南亚）",
    "榜单排名": "罗盘 / 蝉妈妈 / FastMoss / Olive Young",
    "销量": "罗盘 / 蝉妈妈 / FastMoss（Hwahae、Olive Young 不提供）",
    "销售额": "罗盘 / 蝉妈妈 / FastMoss（同上不提供）；海外多币种不折算，留空并写进选品笔记",
    "环比增速": "罗盘 / 蝉妈妈 / FastMoss（同上不提供）",
    "关联达人数": "罗盘 / 蝉妈妈 / FastMoss（同上不提供）",
    "退货率": "仅罗盘 / 蝉妈妈部分类目；FastMoss 不提供；不参与任何规则",
    "核心功效成分": "从商品标题抽出的显性成分词（43 个词表，命中 116 条 / 32%）；标题没写就留空，不推测",
    "剂型": "商详页人工判断",
    "概念标签": "商详页 / 达人话术人工提炼",
    "质地描述": "商详页人工提炼",
    "功效宣称": "商详页可抄，归并需人工",
    "技术壁垒": "内部技术评审",
    "评分": "FastMoss 详情页已自动回填 215 条；蝉妈妈 / Hwahae / Olive Young 也可",
    "评价数": "同上",
    "好评关键词": "抖音罗盘·用户原声（本店 37 个商品已自动取到）",
    "差评关键词": "抖音罗盘·用户原声（本店已接通）；第三方商品需登录天猫 / 小红书 / TikTok 评价后台",
    "评价来源": "按评价实际来源选",
    "备案路径": "NMPA 备案库查询 + 内部法规判断",
    "宣称支撑难度": "内部法规判断",
    "预估成本": "内部成本核算",
    "与我方价格带匹配": "对照自家价格带人工判断",
    "与我方客群匹配": "对照自家客群人工判断",
    "与我方 SKU 重合度": "对照兰本 / 兰至产品线人工判断",
    "决策状态": "内部评审",
    "选品笔记": "人工",
    "数据标记": "系统自动（示例 / 真实）",
}

# ---------- 载入类目树 / 记录 / 成分 ----------
TREE = json.load(open(BUILD + "/data/fm_cat_tree.json", encoding="utf-8"))["tree"]
TREE_ZH = json.load(open(BUILD + "/data/fm_cat_zh.json", encoding="utf-8"))
RECS = json.load(open(BUILD + "/data/records_cat.json", encoding="utf-8"))["records"]


def count_tree(T):
    l2 = sum(len(n.get("children") or []) for n in T)
    l3 = sum(len(g.get("children") or []) for n in T for g in (n.get("children") or []))
    return len(T), l2, l3


N1, N2, N3 = count_tree(TREE)
BEAUTY = [n for n in TREE_ZH["tree"] if (n.get("zh") or "") == "美妆个护"][0]
B2 = BEAUTY.get("children") or []
B3 = sum(len(g.get("children") or []) for g in B2)

c2_cnt = collections.Counter(r.get("二级类目") for r in RECS if r.get("二级类目"))
c3_cnt = collections.Counter(r.get("三级类目") for r in RECS if r.get("三级类目"))
cty_cnt = collections.Counter(r.get("国家/地区") for r in RECS if r.get("国家/地区"))
l1_cnt = collections.Counter(r.get("品类") for r in RECS if r.get("品类"))

# 成分：从「核心功效成分」拆词统计，并取每个成分命中的代表三级类目
ing_cnt = collections.Counter()
ing_cat = collections.defaultdict(collections.Counter)
for r in RECS:
    v = r.get("核心功效成分") or ""
    for w in [s.strip() for s in re.split(r"[、,，/]", v) if s.strip()]:
        ing_cnt[w] += 1
        if r.get("三级类目"):
            ing_cat[w][r["三级类目"]] += 1
ing_cover = sum(1 for r in RECS if r.get("核心功效成分"))

wb = Workbook()

# ================= Sheet 1: 数据源接入台账 =================
ws1 = wb.active
ws1.title = "数据源接入台账"
ws1.append(["数据源", "类型", "覆盖市场", "能拿到", "拿不到", "取数路径", "额度 / 成本", "取数通道", "接入状态"])
SOURCES = [
    ("抖音罗盘", "中国国内平台",
     "中国",
     "销量、销售额、环比增速、关联达人数、一级行业类目、价格、榜单排名、上市日期；用户原声（本店评价）",
     "退货率（仅部分类目）、成分与功效、备案与成本",
     "compass.jinritemai.com → 商品 / 行业 → 榜单页导出",
     "免费（需店铺账号）；行业大盘需额外权限",
     "复用已登录浏览器（本机回环，勿走代理）",
     "列名映射已内置，可直接导入"),
    ("蝉妈妈", "中国国内平台",
     "中国",
     "销量、销售额、环比增速、关联达人数、评分、评价数、价格、榜单排名",
     "三级类目需 AI 判定（榜单不直接给类目树）、成分与功效、备案与成本",
     "chanmama.com → 商品榜 / 品类趋势分析",
     "品牌版品类趋势月报 55 次/月且月度更新，勿按天跑",
     "复用已登录浏览器",
     "列名映射已内置；类目由 ai 判定补齐（50 条里 49 条判定成功）"),
    ("FastMoss", "海外 TikTok Shop 数据",
     "东南亚、欧美、日本",
     "销量、销售额、环比增速、关联达人数、店铺品牌、榜单排名、价格、评分、评价数；"
     "完整三级类目树（30 一级 / 247 二级 / 2370 三级）与国别表",
     "退货率（不提供）、成分与功效、备案与成本",
     "fastmoss.com → Product / Top Selling，region 选站点；类目树走 /api/live/filterInfo",
     "订阅制（标准版 ¥199/月起）；日本站 2025-06 才开，历史浅",
     "SageSurf 跨境浏览器或已登录浏览器（CDP 直连）",
     "已接入，类目树已固化到 _build/data/fm_cat_tree.json"),
    ("Hwahae 화해", "韩国成分与口碑",
     "韩国",
     "核心功效成分、评分、评价数、榜单排名、品类",
     "销量、销售额、环比增速、关联达人数、退货率、备案与成本",
     "hwahae.com → 排名页（无 API、无第三方工具）",
     "无额度限制，但无结构化接口，靠页面解析",
     "普通浏览器即可，需网页采集",
     "新识别，取数脚本待做"),
    ("Olive Young", "韩国渠道榜单",
     "韩国",
     "榜单排名、品类、评分、评价数",
     "销量、销售额、环比增速、关联达人数、退货率、成分与功效",
     "oliveyoung.co.kr 韩国站 → 实时销量榜 / Olive Young Awards",
     "无 API；只公开名次，金额不公开",
     "普通浏览器即可",
     "新识别，取数脚本待做"),
    ("KEV美妆圈", "中国美妆行业数据",
     "中国",
     "品类、细分品类、价格",
     "销量、销售额、环比增速、关联达人数、成分与功效",
     "KEV 中国区数据接口（待打通）",
     "可得性低，接口未通",
     "接口对接中",
     "未打通"),
    ("商品标题（成分抽取）", "降级通道",
     "中国、韩国、日本、欧美、东南亚",
     "核心功效成分（显性成分词）",
     "配方表、浓度、成分顺序",
     "对商品标题做 43 个成分词表匹配（中英文别名 + 词边界）",
     "零成本，但覆盖率低（116/360 = 32%）",
     "脚本 _build/ingredients.py",
     "已在用；标题没写的一律留空，不推测"),
    ("人工调研", "兜底",
     "中国、韩国、日本、欧美、东南亚",
     "与我方 SKU 重合度、与我方价格带匹配、与我方客群匹配、技术壁垒、备案路径、宣称支撑难度、预估成本、决策状态、选品笔记",
     "—",
     "内部评审 / NMPA 备案库查询 / 成本核算",
     "取决于投入，没有平台能替代",
     "人工填写或表格粘贴",
     "已在用（商品档案可手填）"),
]
fill_rows(ws1, SOURCES, colorize=lambda ws, r, row: setattr(
    ws.cell(row=r, column=9), "font",
    OK_FONT if ("已" in str(row[8]) or "内置" in str(row[8])) else (WARN_FONT if ("待" in str(row[8]) or "未" in str(row[8])) else CELL_FONT)))
style_sheet(ws1, [16, 17, 20, 44, 34, 42, 32, 30, 34], 9)
for r in range(2, len(SOURCES) + 2):
    ws1.row_dimensions[r].height = 92

# ================= Sheet 2: 市场可走通性 =================
ws2 = wb.create_sheet("市场可走通性")
ws2.append(["市场", "可用数据源", "四件套是否齐全", "结论", "说明", "卡点 / 下一步", "本库实际国家/地区"])
MK = [
    ("中国", "抖音罗盘、蝉妈妈、KEV美妆圈（未通）", "齐全",
     "可走通",
     "罗盘与蝉妈妈都能供销量、销售额、环比增速、关联达人数，链路完整。注意口径是「抖音电商」，与 TikTok Shop 系市场不可直接横向比较。",
     "KEV 接口未打通，暂缺行业侧品类趋势；罗盘行业大盘需额外权限。",
     "中国 110 条（另有 10 条示例）"),
    ("韩国", "Hwahae 화해、Olive Young", "缺失",
     "走不通",
     "FastMoss 不覆盖韩国（TikTok Shop 韩国站尚未开通）。目前只能拿到成分、评分、评价数与榜单名次，拿不到销量与销售额。",
     "「窗口期」「改良机会」「风险提醒」「空白赛道」四条规则在韩国数据上算不出来；需为韩国单独做 Hwahae / Olive Young 的网页采集脚本，或直接接受韩国只做成分情报。",
     "0 条真实（7 条为示例数据，标了 Hwahae / Olive Young）"),
    ("日本", "FastMoss（日本站）", "齐全但浅",
     "勉强能用",
     "FastMoss 有日本站，四件套齐全；但 2025-06 才开站，历史数据浅，环比基数不稳。",
     "环比建议改看 90 天或绝对销量，别只看环比增速；后续可补 Qoo10 / 乐天榜单做交叉验证。",
     "日本 1 条真实（6 条为示例）"),
    ("欧美", "FastMoss（US / GB / DE / FR / IT / ES / MX / BR）", "齐全",
     "可走通",
     "FastMoss 覆盖美英德法意西墨巴，四件套完整，是当前质量最好的海外数据源。",
     "退货率 FastMoss 不提供，改良机会类结论需要另找评论源或人工估算；海外价多币种不折算，销售额字段留空。",
     "美国 14 / 巴西 6 / 墨西哥 2 / 英国 1 条"),
    ("东南亚", "FastMoss（ID / VN / TH / MY / PH / SG）", "齐全",
     "可走通",
     "FastMoss 在东南亚覆盖最好、历史最长，四件套完整，直播带货数据尤其完整——本库 226/360 条来自这里。",
     "六个站点口径不同，跨站点汇总前先确认统计周期一致。",
     "印尼 73 / 泰国 63 / 越南 42 / 菲律宾 27 / 马来西亚 20 / 新加坡 1 条"),
]
fill_rows(ws2, MK, colorize=lambda ws, r, row: setattr(
    ws.cell(row=r, column=4), "font", OK_FONT if row[3] == "可走通" else (WARN_FONT if row[3] == "走不通" else CELL_FONT)))
style_sheet(ws2, [10, 34, 15, 12, 62, 48, 34], 7)
for r in range(2, len(MK) + 2):
    ws2.row_dimensions[r].height = 92

# ================= Sheet 3: 洞察规则与前提（5 → 9 条） =================
ws3 = wb.create_sheet("洞察规则与前提")
ws3.append(["分组", "规则", "判定条件", "依赖字段", "其中平台可导", "其中需人工标", "能否自动算出", "说明"])
RULES = [
    ("原有 5 条", "窗口期", "环比增速 ≥ 60%（达人数与销售额不设门槛）",
     "环比增速", "环比增速", "无", "可自动",
     "只看增速。真实 TikTok 上升品的关联达人数中位约 110，旧阈值「≤15」是示例数据量纲，在真实数据上命中 0，因此取消达人数与销售额门槛。"
     "关联达人数改为展示用的「达人渗透度」：≤50 未铺开 / ≤150 起量 / >150 已铺开。"),
    ("原有 5 条", "改良机会", "差评关键词非空即命中",
     "差评关键词", "无", "差评关键词（来自商品评价）", "需先补评价数据",
     "原口径「退货率 ≥ 10% 且 销量 ≥ 3 万」已废弃——三个平台的商品榜单都不提供退货率，抓不到的字段不能当规则前提。"
     "改为评价驱动：差评关键词不是缺点清单，是改良清单。"),
    ("原有 5 条", "空白赛道", "与我方 SKU 重合度 = 全新 且 环比增速 ≥ 25",
     "与我方 SKU 重合度、环比增速", "环比增速", "与我方 SKU 重合度", "不能",
     "「与我方 SKU 重合度」任何平台都导不出来，只能对照兰本 / 兰至产品线人工标。不标这个字段，这条规则永远命中 0 —— 这不是「没机会」，是「还没标」。"),
    ("原有 5 条", "可落地", "备案路径 = 普通化妆品备案 且 宣称支撑难度 = 无需评价 且 预估成本 ÷ 价格 ≤ 25%",
     "备案路径、宣称支撑难度、预估成本、价格", "价格", "备案路径、宣称支撑难度、预估成本", "不能",
     "四个字段里三个是内部判断，纯人工。好处是不依赖任何平台数据，所以连韩国数据也能参与这条规则的判定。"),
    ("原有 5 条", "风险提醒", "环比增速 < 0 或 差评关键词含高危词",
     "环比增速、差评关键词", "环比增速", "差评关键词", "基本可自动",
     "高危词表：过敏 / 刺激 / 烂脸 / 闷痘 / 红肿 / 刺痛 / 泛红 / 发痒 等。负增长本身已足以触发提醒，不依赖退货率。"),
    ("v1.5 新增", "市场机会", "某「国家/地区 × 二级类目」里 ≥3 个商品 且 平均环比增速 ≥ 60%",
     "国家/地区、二级类目、环比增速、销量", "全部（FastMoss region 直出）", "无", "可自动",
     "把「哪个国家的哪个功能子类正在起来」直接摆出来。数据支持最好的一条规则——国家与类目都是平台原生字段，不依赖任何人工标注。"
     "门槛用条数而不是销售额，因为海外记录销售额多为空（多币种不折算）。"),
    ("v1.5 新增", "增速品类", "单个二级类目 ≥5 个商品 且 平均环比增速 ≥ 80% 且 榜内增速中位 > 0",
     "二级类目、环比增速", "全部", "无", "可自动",
     "看的是类目整体起飞而不是单个爆品，用来判断「这个功能子类值不值得开一条产品线」。"
     "要求中位数为正，是为了避免被一两个异常增速值拉高平均值。"),
    ("v1.5 新增", "新品机会", "上市日期距今 ≤ 90 天 且 环比增速 ≥ 100%",
     "上市日期、环比增速", "上市日期、环比增速", "无", "可自动",
     "新品能同时拿到「刚上市」和「已经在涨」，说明产品定义本身对路，而不是靠时间积累。"
     "中国 / 海外都能算，是判断「这个新品逻辑能不能抄」最快的一条。"),
    ("v1.5 新增", "成分机会", "同一成分被 ≥3 个商品使用 且 覆盖 ≥2 个不同类目 且 平均环比增速 ≥ 50%",
     "核心功效成分、三级类目、环比增速", "环比增速", "核心功效成分（标题抽取，32% 覆盖）", "部分可算",
     "成分是这个平台给自家品牌选方向的落点：一个成分跨类目反复出现，说明它是「通用解法」而不是某条产品线的专属卖点。"
     "注意成分是从标题抽的显性词，覆盖只有 32%，命中为空要区分「没人用」和「标题没写」。"),
]
fill_rows(ws3, RULES, colorize=lambda ws, r, row: setattr(
    ws.cell(row=r, column=7), "font", OK_FONT if str(row[6]).startswith("可自动") or str(row[6]).startswith("基本") else WARN_FONT))
style_sheet(ws3, [11, 12, 44, 30, 24, 28, 14, 70], 8)
for r in range(2, len(RULES) + 2):
    ws3.row_dimensions[r].height = 96

# ================= Sheet 4: 字段可得性 =================
ws4 = wb.create_sheet("字段可得性")
ws4.append(["序号", "层级", "字段名", "类型", "可得性", "取值 / 来源", "可选值"])
rows, i, last_g = [], 0, None
for f in FIELDS:
    if f["g"] != last_g:
        rows.append(("###" + f["g"], "", "", "", "", "", ""))
        last_g = f["g"]
    i += 1
    rows.append((str(i), f["g"], f["n"], f["t"], SRC_LABEL.get(f["src"], f["src"]),
                 SRC_OWNER.get(f["n"], ""), " / ".join(f["opts"]) if f["opts"] else ""))
fill_rows(ws4, rows, group_col=1, colorize=lambda ws, r, row: setattr(
    ws.cell(row=r, column=5), "font",
    OK_FONT if row[4] == "平台可导" else (CELL_FONT if row[4] == "部分可导" else WARN_FONT)))
style_sheet(ws4, [6, 14, 20, 10, 12, 56, 34], 7)
auto = sum(1 for f in FIELDS if f["src"] == "auto")
half = sum(1 for f in FIELDS if f["src"] == "half")
man = sum(1 for f in FIELDS if f["src"] == "manual")
r = len(rows) + 3
ws4.cell(row=r, column=1, value="统计").font = GRP_FONT
ws4.cell(row=r, column=2, value="平台可导 %d 个 / 部分可导 %d 个 / 需人工标 %d 个 / 合计 %d 个" % (auto, half, man, len(FIELDS))).font = CELL_FONT
ws4.cell(row=r + 1, column=1, value="结论").font = GRP_FONT
ws4.cell(row=r + 1, column=2,
         value="平台能自动填的字段不到一半。导入数据后，凡是依赖人工字段的洞察规则都不会自动出结果 —— "
               "引擎会在「机会洞察」页明确提示缺哪些字段，避免把「还没标」误读成「没机会」。").font = CELL_FONT

# ================= Sheet 5: 品类体系（v1.5 新增） =================
ws5 = wb.create_sheet("品类体系")
ws5.append(["分组", "层级", "名称", "类目ID / 说明", "上属", "本库记录数", "备注"])
cat_rows = []

cat_rows.append(("###A · 平台类目树（FastMoss /api/live/filterInfo）", "", "", "", "", "", ""))
cat_rows.append(("平台类目树", "总览", "完整三级类目树", "30 / 247 / 2370",
                 "一级 / 二级 / 三级", str(len(RECS)), "接口无需签名，同源 fetch 直连即可取到；已固化到 _build/data/fm_cat_tree.json"))
cat_rows.append(("平台类目树", "一级", "美妆个护", "cid 14",
                 "—", str(l1_cnt.get("个护", 0) + l1_cnt.get("护肤", 0) + l1_cnt.get("彩妆", 0) + l1_cnt.get("身体", 0) + l1_cnt.get("香氛", 0) + l1_cnt.get("工具", 0)),
                 "本库只用这一条分支；其余 29 个一级（服装 / 3C / 家居…）与美妆业务无关，未译名"))
cat_rows.append(("平台类目树", "二级", "共 %d 个（美妆个护下）" % len(B2), "—", "美妆个护", "",
                 "如 手足及指甲护理 / 洗浴与身体护理 / 美容护肤 / 头部护理与造型 / 美容、个护电器"))
for g in B2:
    kids = g.get("children") or []
    cat_rows.append(("平台类目树", "三级归属", g.get("zh") or g.get("name"), "cid " + str(g.get("cid")),
                     "美妆个护", "", "%d 个三级类目" % len(kids)))
cat_rows.append(("平台类目树", "三级", "共 %d 个（美妆个护下）" % B3, "—", "美妆个护", "",
                 "本库实际用到的三级类目只占其中一部分（见分组 C）"))

cat_rows.append(("###B · 功能子类（页面「二级类目」，由三级归并）", "", "", "", "", "", ""))
for k, v in sorted(c2_cnt.items(), key=lambda x: -x[1]):
    cat_rows.append(("功能子类", "二级", k, "—", "由三级类目归并", str(v),
                     "归并规则见 _build/map_cats.py 的 FUNC_OF_L3"))
cat_rows.append(("功能子类", "合计", "%d 个功能子类" % len(c2_cnt), "—", "—", str(sum(c2_cnt.values())),
                 "一个功能子类可对应平台里的多个三级类目（如 身体清洁 含 沐浴露与香皂 + 男士洗浴与身体护理）"))

cat_rows.append(("###C · 三级类目分布（页面「三级类目」，平台最细一级）", "", "", "", "", "", ""))
for k, v in sorted(c3_cnt.items(), key=lambda x: -x[1]):
    cat_rows.append(("三级类目", "三级", k, "—", "平台三级", str(v), ""))
cat_rows.append(("三级类目", "合计", "%d 个三级类目" % len(c3_cnt), "—", "—", str(sum(c3_cnt.values())),
                 "360 条真实记录 100% 映射成功；示例记录另按显式映射表补齐"))

cat_rows.append(("###D · 映射方法（三级类目从哪来）", "", "", "", "", "", ""))
cat_rows.append(("映射方法", "通道 1", "FastMoss 原生类目路径", "占 250/360 条", "榜单元数据 all_category_name",
                 "250", "最权威：平台自己给的三级路径，直接采信"))
cat_rows.append(("映射方法", "通道 2", "抖音罗盘行业类目", "占 60/360 条", "罗盘榜单行业字段",
                 "60", "罗盘给一级行业类目，按 平台二级 → 业务大类 映射表归位"))
cat_rows.append(("映射方法", "通道 3", "Kimi 判定", "占 49/360 条", "蝉妈妈榜单",
                 "49", "蝉妈妈不直接给类目树；限定候选清单让模型挑，挑不出给「未定」，50 条里 49 条成功"))
cat_rows.append(("映射方法", "通道 4", "手工补", "占 1/360 条", "—", "1", "末条（温博士 B5 水）人工补 爽肤水、化妆水"))
cat_rows.append(("映射方法", "校验", "用平台真实中文路径反校验 AI 译名", "66 / 66 命中", "—", "",
                 "从榜单里捞到 66 条平台原生的中文三级路径，反向校验 AI 英文译名，全部对齐后才敢入库"))
fill_rows(ws5, cat_rows, group_col=1)
style_sheet(ws5, [12, 11, 30, 24, 20, 12, 62], 7)

# ================= Sheet 6: 成分库（v1.5 新增） =================
ws6 = wb.create_sheet("成分库")
ws6.append(["序号", "成分", "命中商品数", "覆盖类目数", "代表类目", "说明"])
ing_rows = []
for idx, (k, v) in enumerate(ing_cnt.most_common(), 1):
    cats = ing_cat.get(k) or collections.Counter()
    ing_rows.append((str(idx), k, str(v), str(len(cats)),
                     " / ".join(c for c, _ in cats.most_common(3)) or "—",
                     "标题里出现该成分词的商品数"))
fill_rows(ws6, ing_rows, colorize=lambda ws, r, row: setattr(
    ws.cell(row=r, column=3), "font", OK_FONT if int(row[2]) >= 5 else CELL_FONT))
style_sheet(ws6, [6, 18, 12, 11, 40, 40], 6)
r = len(ing_rows) + 3
for label, val in [
    ("统计", "%d 个成分词 / 命中 %d 条记录（覆盖率 %.0f%%）/ 词表见 _build/ingredients.py" % (
        len(ing_cnt), ing_cover, 100.0 * ing_cover / len(RECS))),
    ("抽取规则", "只从商品标题抽显性成分词，中英文别名 + 词边界匹配；标题没写就留空，不做任何推测。"),
    ("为什么降级", "五条详情页通道全部不可用，详见下方「成分通道排查」说明"),
    ("怎么读", "命中数高 = 平台上的通用解法（如香精 42 个品、烟酰胺 20 个品）；"
             "覆盖类目数多 = 跨品类的通用卖点，对自家开新品线更有参考价值。"),
    ("局限", "覆盖率仅 32%：命中为空要先区分「没人用这个成分」和「卖家标题没写」。"
             "要看真实配方表，必须走详情页通道（FastMoss 详情页游客态字段全空、抖音 haohuo 需 APP 扫码、"
             "淘宝天猫小红书 TikTok 浏览器均未登录、国家药监局备案库返回 412 反爬）。"),
]:
    ws6.cell(row=r, column=1, value=label).font = GRP_FONT
    ws6.cell(row=r, column=2, value=val).font = CELL_FONT
    r += 1

# ================= Sheet 7: 三平台取数映射 =================
ws7 = wb.create_sheet("三平台取数映射")
ws7.append(["工作台字段", "抖音罗盘列名", "蝉妈妈列名", "FastMoss 列名", "备注"])
MAPPING = [
    ("商品名称", "商品名称 / 商品标题", "商品名称 / 商品标题", "product_name / title", "导入按此字段去重"),
    ("品牌", "品牌 / 店铺名称", "品牌 / 店铺", "shopname / brand", ""),
    ("商品ID", "商品ID", "商品ID", "product_id", "可留空"),
    ("价格", "到手价 / 客单价", "到手价 / 均价", "price", "海外价折算见版本说明待定项 1"),
    ("销量", "销量 / 销售件数", "销量 / 近30天销量", "units_sold / total_units_sold", "「万 / k / 千」自动换算"),
    ("销售额", "成交额 / 成交金额", "销售额", "gmv / total_gmv", "多币种不折算，海外记录留空并写进选品笔记"),
    ("环比增速", "环比增速", "环比增速 / 增长率", "growth_rate", "取数周期见版本说明待定项 2"),
    ("关联达人数", "关联达人数 / 带货达人数", "关联达人数", "—（FastMoss 商品页给达人列表，需计数）", "FastMoss 需二次统计"),
    ("退货率", "退货率（部分类目）", "退货率 / 退款率", "—（三个平台商品榜都不提供）", "不参与任何规则，仅备查"),
    ("国家/地区", "—（国内固定中国）", "—（国内固定中国）", "region（US / GB / DE / FR / IT / ES / MX / BR / ID / VN / TH / MY / PH / SG / JP + 欧陆小站）",
     "比「所属市场」大区细一级；韩国由 Hwahae / Olive Young 人工补"),
    ("二级类目", "一级行业类目 → 映射表归位", "由 AI 判定后归位", "all_category_name 的二级段 + FUNC_OF_L3 归并",
     "46 种功能子类，是「品类细分树」的中间层"),
    ("三级类目", "—（不直接给）", "—（不直接给）", "all_category_name 的三级段（平台最细类目）",
     "71 种；FastMoss 是唯一直接给三级类目的源"),
    ("类目ID", "—", "—", "cid 链，如 14/848648/601554", "按它可精确复现取数，不受译名影响"),
    ("榜单排名", "榜单 / 排名", "排名 / 榜单", "rank", "带上平台与类目一起记"),
    ("所属市场", "—（国内固定中国）", "—（国内固定中国）", "region 归大区", "中国 / 韩国 / 日本 / 欧美 / 东南亚"),
    ("上市日期", "上架时间 / 发布时间", "上架时间", "launch_date", "「新品机会」规则的前提"),
    ("品类", "一级类目", "一级类目", "category", "归并成 6 个业务大类"),
    ("细分品类", "二级 / 三级类目", "二级 / 三级类目", "subcategory", "v1.5 前的老字段，保留兼容"),
    ("核心功效成分", "—（详情页，未打通）", "—", "—", "改从商品标题抽取，覆盖 32%"),
    ("数据来源", "—（导入时自动打标）", "—（导入时自动打标）", "—（导入时自动打标）", "在导入第 1 步选源"),
]
fill_rows(ws7, MAPPING)
style_sheet(ws7, [16, 26, 24, 44, 40], 5)
for r in range(2, len(MAPPING) + 2):
    ws7.row_dimensions[r].height = 34

# ================= Sheet 8: 自有店铺口碑库 =================
ws8 = wb.create_sheet("自有店铺口碑库")
ws8.append(["字段", "类型", "取数来源", "口径说明"])
REVIEW_FIELDS = [
    ("商品名称", "text", "罗盘·用户原声·商品明细", "本店在售商品名"),
    ("商品ID", "text", "同上", "与「全球新品库」的商品ID同体系（抖店 product_id）"),
    ("店铺类目", "text", "同上", "罗盘给的三级类目，如 个人护理/身体护理/身体乳/霜/贴/膏/油"),
    ("评价数", "number", "同上", "统计周期内累计评价条数"),
    ("好评数", "number", "同上", "好评条数"),
    ("好评率", "number", "同上", "好评数 / 评价数 × 100，存百分比数值（如 93.04）"),
    ("差评订单数", "number", "同上", "产生差评标签的订单数"),
    ("差评率", "number", "同上", "商品差评订单数 / 物流签收订单数 × 100（罗盘口径，分母不是评价数）"),
    ("评价差评率", "number", "同上", "差评评价 / 评价数 × 100，与上一行分母不同，别混用"),
    ("好评关键词", "text", "罗盘·用户原声·comment_type=1", "如「物美价廉 / 推荐 / 会回购」，多条用「 / 」分隔"),
    ("差评关键词", "text", "罗盘·用户原声·comment_type=2", "如「描述不符 / 不推荐」"),
    ("差评原因", "text", "罗盘·用户原声·comment_type=4", "平台打的差评标签+条数，如「虚假宣传 4 / 效果虚假 2」——改良清单就从这来"),
    ("差评原声", "text", "同上", "买家差评原文样例，几条用「 | 」分隔，只截 900 字"),
    ("品质退货数", "number", "罗盘·商品明细", "商品品质退货订单量"),
    ("品质退货率", "number", "同上", "品质退货 / 签收 × 100"),
    ("投诉数", "number", "同上", "投诉工单量"),
    ("投诉率", "number", "同上", "投诉 / 签收 × 100"),
    ("采集周期", "text", "—", "如 2026/08/20 ~ 2026/09/18（近 30 天）"),
    ("数据来源", "select", "—", "固定「抖音罗盘」"),
    ("评价来源", "select", "—", "固定「抖音评价」"),
    ("备注", "text", "—", "通道说明：抖音电商罗盘 · 体验 · 用户原声（本店）"),
]
fill_rows(ws8, REVIEW_FIELDS)
style_sheet(ws8, [16, 10, 34, 64], 4)
for r in range(2, len(REVIEW_FIELDS) + 2):
    ws8.row_dimensions[r].height = 30

# ================= Sheet 9: 版本说明 =================
ws9 = wb.create_sheet("版本说明")
ws9.append(["项目", "内容"])
NOTES = [
    ("版本", "v1.5 —— 与已上线「全球选品平台」（线上 v8）对齐；新增国家市场维度、三级品类细分、成分库，洞察规则 5 条 → 9 条"),
    ("v1.5 改了什么",
     "① 字段 35 → 39：新增「二级类目 / 三级类目 / 类目ID / 国家/地区」，360 条真实记录 100% 映射类目、383 条（含示例）落到具体国家；"
     "② 看板新增「国家/地区市场」六列对比表与「二级类目增速榜」；"
     "③ 品类库视图改为双 Tab：品类细分树（可逐层展开）+ 成分库；"
     "④ 成分库 43 个成分，按命中商品数排行；"
     "⑤ 新增 4 条机会洞察：市场机会 / 增速品类 / 新品机会 / 成分机会"),
    ("v1.5 为什么改",
     "原来只有「一级品类 + 大区」，看不出「哪个国家的哪个功能子类在涨」——彩妆里底妆和眼妆可能是相反走势，"
     "护肤里的保湿乳和洁面也不是一回事。国家与三级类目都是平台原生字段，不加人工成本就能拿到，"
     "而这两维恰好是选方向时最先要看的两个坐标。"),
    ("v1.4 改了什么", "口碑数据真正落库：抖音后台「用户原声」37 个本店商品 → 自有店铺口碑库；FastMoss 详情页评分/评价数 215 条回填；360 条记录 AI 增强（中文品名 + AI 速读）"),
    ("v1.4 为什么改", "「改良机会」规则需要差评关键词驱动，但之前评价字段全空；抖音后台是唯一不用额外登录就能拿到评价原文的通道"),
    ("v1.3 改了什么", "窗口期取消达人数与销售额门槛（只看环比增速）；改良机会由退货率驱动改为差评关键词驱动；风险提醒改为「环比 < 0 或差评含高危词」；口碑层新增「好评关键词」「评价来源」两个字段"),
    ("v1.3 为什么改", "旧阈值「关联达人数 ≤ 15」来自示例数据量纲，真实数据中位 110 → 规则命中 0，等于空转；退货率三个平台榜单都不提供 → 不能拿抓不到的字段当前提"),
    ("v1.2 改了什么", "新增：数据源接入台账 / 市场可走通性 / 洞察规则与前提 / 字段可得性 四张表；字段清单改为从页面自动抽取，不再手工维护"),
    ("核心结论 1", "没有单一数据源能覆盖全部五个市场：罗盘 / 蝉妈妈只覆盖中国，FastMoss 覆盖东南亚 / 欧美 / 日本，韩国两家都不覆盖"),
    ("核心结论 2", "韩国是唯一走不通的市场：TikTok Shop 韩国站未开通，FastMoss 无数据；Hwahae / Olive Young 只给成分口碑与名次，给不了销量销售额"),
    ("核心结论 3", "%d 个字段里平台能自动供的只有 %d 个，另有 %d 个部分可导、%d 个任何源都给不了。9 条洞察规则里有 3 条完全依赖人工字段，导入数据不会自动出结果" % (
        len(FIELDS), auto, half, man)),
    ("核心结论 4", "品类细分只有 FastMoss 能直出到三级（30 / 247 / 2370 的完整类目树，且带 cid）；罗盘给一级行业类目、蝉妈妈干脆不给——"
                "所以「类目」这件事要按源分通道处理，不能指望一套映射打天下"),
    ("核心结论 5", "完整配方表目前拿不到：五条详情页通道（FastMoss 详情页游客态字段全空、抖音 haohuo 需 APP 扫码、"
                "淘宝/天猫/小红书/TikTok 浏览器未登录、国家药监局备案库 412 反爬）全部不可用，成分只能从标题抽显性词，覆盖 32%"),
    ("字段数", "选品库 %d 个（v1.4 的 35 个 + 二级类目 / 三级类目 / 类目ID / 国家/地区）；另建「自有店铺口碑库」21 个字段" % len(FIELDS)),
    ("数据来源可选值", "抖音罗盘 / 蝉妈妈 / FastMoss / Hwahae / Olive Young / KEV美妆圈 / 人工调研 / 其他"),
    ("待定项 1（待拍板）", "海外价折算汇率与含税口径：建议统一按月初中间价折算、取含税到手价（当前海外销售额一律留空）"),
    ("待定项 2（待拍板）", "环比统计周期取 30 天还是 90 天、大促月是否剔除假峰值：建议默认 30 天并人工剔除大促；日本站因历史浅建议看 90 天"),
    ("待定项 3（待拍板）", "「与我方 SKU 重合度」基准产品线：兰本 / 兰至 / 两条线合并判断"),
    ("待定项 4（待拍板）", "韩国是否值得单独做 Hwahae / Olive Young 的网页采集脚本，还是接受韩国只做成分口碑情报"),
    ("待定项 5（评价通道）", "抖音后台已走通（只覆盖本店商品）；天猫 / 淘宝 / 小红书 / TikTok 在 SageSurf 与 webbridge 两台浏览器上都未登录，"
                        "需登录后照 _build/wb_douyin_usersound.py 的写法各接一条取数脚本"),
    ("待定项 6（成分通道）", "要看完整成分表，需登录任一天猫 / 抖音商城 / 小红书 / TikTok 的商详页；"
                        "或改用可视化成分数据库（如美丽修行）做人工对照，但要接受它是第三方整理而非官方配方"),
    ("环境提示", "本表为静态口径文档，不含任何 Excel 公式，可在任意设备打开"),
]
fill_rows(ws9, NOTES)
style_sheet(ws9, [22, 112], 2)
for r in range(2, len(NOTES) + 2):
    ws9.row_dimensions[r].height = 44

wb.save(OUT)
print("SAVED", OUT, os.path.getsize(OUT))
print("字段统计：auto %d / half %d / manual %d / 合计 %d" % (auto, half, man, len(FIELDS)))
print("类目：二级 %d 种 / 三级 %d 种 | 成分 %d 个 / 覆盖 %d 条" % (len(c2_cnt), len(c3_cnt), len(ing_cnt), ing_cover))
print("Sheet:", wb.sheetnames)

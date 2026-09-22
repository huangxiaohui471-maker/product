try:
    import openpyxl
except ImportError:
    import subprocess, sys
    subprocess.check_call([sys.executable, "-m", "pip", "install", "--quiet", "openpyxl>=3.1.0"])
    import openpyxl

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import CellIsRule


def xl_color(css_hex: str) -> str:
    value = css_hex.removeprefix("#").upper()
    if len(value) != 6:
        raise ValueError(f"Expected #RRGGBB, got: {css_hex}")
    return "FF" + value


XL_HEADER_BG = xl_color("#4472C4")
XL_HEADER_FG = xl_color("#FFFFFF")
XL_LAYER_BG = xl_color("#EDF2FB")
XL_INPUT_BG = xl_color("#FFFBF0")
XL_TOTAL_BG = xl_color("#2F5597")
XL_BORDER = xl_color("#BFBFBF")
XL_RED_BG = xl_color("#FFC7CE")
XL_RED_FG = xl_color("#9C0006")
XL_YEL_BG = xl_color("#FFEB9C")
XL_YEL_FG = xl_color("#9C6500")

THIN = Side(style="thin", color=XL_BORDER)
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

F_HEAD = Font(size=11, bold=True, color=XL_HEADER_FG)
F_BODY = Font(size=10)
F_TOTAL = Font(size=11, bold=True, color=XL_HEADER_FG)

AL_C = Alignment(horizontal="center", vertical="center")
AL_L = Alignment(horizontal="left", vertical="center", wrap_text=True)

wb = Workbook()
wb.properties.title = "全球新品选品中台 · P1 字段口径表"


def write_head(ws, headers):
    for c, name in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=c, value=name)
        cell.font = F_HEAD
        cell.fill = PatternFill("solid", fgColor=XL_HEADER_BG)
        cell.alignment = AL_C
        cell.border = BORDER
    ws.row_dimensions[1].height = 24


def style_body(ws, first_row, last_row, ncol, wrap_cols=(), center_cols=(), fill_map=None):
    fill_map = fill_map or {}
    for r in range(first_row, last_row + 1):
        for c in range(1, ncol + 1):
            cell = ws.cell(row=r, column=c)
            cell.font = F_BODY
            cell.border = BORDER
            if c in center_cols:
                cell.alignment = AL_C
            elif c in wrap_cols:
                cell.alignment = AL_L
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")
            if c in fill_map:
                cell.fill = PatternFill("solid", fgColor=fill_map[c])


# ============ Sheet 1: 字段口径表 ============
ws1 = wb.active
ws1.title = "字段口径表"
write_head(ws1, ["序号", "字段名", "所属层", "类型", "口径定义", "取值示例", "必填", "数据来源", "可得性", "评审意见"])

fields = [
    (1, "商品名称", "①商品主数据", "文本", "用品牌官方命名，不用平台改写版或达人叫法", "「XX 屏障修护精华油」", "是", "各平台商品页", "人工录入"),
    (2, "品牌", "①商品主数据", "文本", "品牌全称，同一品牌在全表统一写法", "「Dr.Jart+」", "是", "各平台商品页", "人工录入"),
    (3, "所属系列", "①商品主数据", "文本", "品牌内部产品线；没有归属填「无」", "「水动力系列」", "否", "品牌官网", "人工录入"),
    (4, "品类", "①商品主数据", "枚举", "护肤 / 彩妆 / 个护 / 身体 / 香氛 / 工具", "身体", "是", "推断", "人工录入"),
    (5, "细分品类", "①商品主数据", "文本", "比品类再细一档，用行业通行叫法", "「身体冷霜」", "是", "推断", "人工录入"),
    (6, "规格", "①商品主数据", "文本", "数字+单位，中间不加空格", "「50g」", "是", "商品页", "人工录入"),
    (7, "价格", "①商品主数据", "货币（¥）", "取到手价（含券后），不用标价；海外价按统一汇率折算", "269.00", "是", "商品页", "人工录入"),
    (8, "上市日期", "①商品主数据", "日期", "首次上架日；找不到就填首次被收录的日期", "2026-08-15", "是", "商品页", "人工录入"),
    (9, "主战场渠道", "①商品主数据", "枚举", "抖音 / 天猫 / 小红书 / 亚马逊 / TikTok Shop / 线下", "抖音", "是", "榜单数据", "批量可导"),
    (10, "备案号", "①商品主数据", "文本", "中国填 NMPA 备案号；海外品填当地编号，没有填「无」", "「国妆网备进字…」", "否", "药监局备案库", "自动可抓"),

    (11, "核心功效成分", "②成分功效", "文本", "只填主打成分，最多 3 个，用顿号分隔", "「神经酰胺、角鲨烷」", "是", "备案成分表", "人工录入"),
    (12, "成分浓度档", "②成分功效", "枚举", "低（<1%）/ 中（1%–5%）/ 高（>5%）/ 未公开", "中", "否", "推断", "人工录入"),
    (13, "全成分表", "②成分功效", "文本", "按 INCI 名称完整录入，顺序不可调整", "「水、甘油、角鲨烷…」", "否", "备案信息", "自动可抓"),
    (14, "剂型", "②成分功效", "枚举", "水 / 乳 / 霜 / 油 / 膏 / 喷雾 / 粉", "油", "是", "商品页", "人工录入"),
    (15, "质地描述", "②成分功效", "文本", "保留原话描述肤感，不要改写成专业词", "「冷霜质地、夏天也敢用」", "否", "商品页 / UGC", "人工录入"),
    (16, "功效宣称", "②成分功效", "文本", "官方宣称，最多 3 条", "「屏障修护」", "是", "商品页", "人工录入"),
    (17, "概念标签", "②成分功效", "文本", "行业概念词，最多 3 个", "「以油养肤、早C晚A」", "是", "内容平台", "人工录入"),
    (18, "技术壁垒", "②成分功效", "枚举", "无壁垒 / 配方工艺 / 独家原料 / 专利技术", "配方工艺", "是", "推断", "人工录入"),
    (19, "关键原料可得性", "②成分功效", "枚举", "易得 / 需进口 / 独家供应 / 受限", "需进口", "否", "供应商询价", "人工录入"),
    (20, "与我方 SKU 重合度", "②成分功效", "枚举", "全新 / 部分重合 / 高度重合", "部分重合", "是", "我方产品库比对", "人工录入"),

    (21, "榜单排名", "③市场表现", "文本", "写清「平台+榜单+名次」，多个榜单用分号隔开", "「抖音身体护理第 12」", "否", "平台榜单", "批量可导"),
    (22, "环比增速", "③市场表现", "百分比", "近 30 天销售额 ÷ 前 30 天销售额 − 1", "0.42（即 +42%）", "是", "工具数据", "批量可导"),
    (23, "价格带", "③市场表现", "枚举", "<100 / 100–300 / 300–600 / 600–1000 / >1000", "100–300", "是", "由价格推导", "批量可导"),
    (24, "主销渠道占比", "③市场表现", "文本", "各渠道销售额占比，只取前两位", "「抖音 70%；天猫 30%」", "否", "工具数据", "批量可导"),
    (25, "所属市场", "③市场表现", "枚举", "中国 / 韩国 / 日本 / 欧美 / 东南亚", "韩国", "是", "来源站点", "人工录入"),
    (26, "生命周期阶段", "③市场表现", "枚举", "萌芽 / 爬坡 / 成熟 / 衰退", "爬坡", "是", "由增速判断", "人工录入"),

    (27, "评价数", "④口碑内容", "数字", "该商品主链接的累计评价 / 评论数", "12400", "是", "平台商品页", "批量可导"),
    (28, "评分", "④口碑内容", "数字（1位小数）", "5 分制；其他平台先折算到 5 分再录", "4.6", "是", "平台商品页", "批量可导"),
    (29, "好评关键词", "④口碑内容", "文本", "从好评里提炼，最多 3 个高频词", "「不黏、好吸收」", "否", "UGC 评论", "人工录入"),
    (30, "差评关键词", "④口碑内容", "文本", "从差评里提炼，最多 3 个高频词——这是开发机会的直接线索", "「闷痘、味道冲」", "是", "UGC 评论", "人工录入"),
    (31, "热议场景", "④口碑内容", "文本", "UGC 里反复出现的使用场景", "「浴后 3 分钟内」", "否", "UGC 内容", "人工录入"),
    (32, "达人带货密度", "④口碑内容", "枚举", "无 / 零星 / 中等 / 铺天盖地", "中等", "否", "内容平台", "人工录入"),

    (33, "备案路径", "⑤可行性", "枚举", "普通化妆品备案 / 特殊化妆品注册 / 进口备案 / 暂无", "普通化妆品备案", "是", "NMPA", "自动可抓"),
    (34, "特殊化妆品风险", "⑤可行性", "枚举", "无 / 需美白 / 需防晒 / 需染发烫发 / 需祛斑", "无", "是", "NMPA 分类", "自动可抓"),
    (35, "功效宣称支撑难度", "⑤可行性", "枚举", "无需评价 / 需文献资料 / 需人体功效试验", "需文献资料", "是", "法规知识库", "人工录入"),
    (36, "推荐代工", "⑤可行性", "文本", "能做该剂型的代工厂，可写多家", "「XX 生物」", "否", "供应商库", "人工录入"),
    (37, "预估成本带", "⑤可行性", "货币（¥）", "单件出厂成本区间的中值", "42.00", "否", "供应商询价", "人工录入"),
    (38, "起订量 MOQ", "⑤可行性", "数字", "按单件计，不带千分位以外的说明", "10000", "否", "供应商询价", "人工录入"),
    (39, "预估交期（天）", "⑤可行性", "数字", "从打样到出货的自然日", "75", "否", "供应商询价", "人工录入"),
    (40, "与我方价格带匹配", "⑤可行性", "枚举", "匹配 / 偏高 / 偏低", "匹配", "是", "我方产品库", "人工录入"),
    (41, "与我方客群匹配", "⑤可行性", "枚举", "匹配 / 需教育 / 不符", "匹配", "是", "我方人群数据", "人工录入"),

    (42, "综合得分", "⑥决策", "数字（1位小数）", "按「评分卡规则」7 个维度加权算出，0–100 分", "78.5", "否", "评分卡", "自动可抓"),
    (43, "决策状态", "⑥决策", "枚举", "待评 / 候选池 / 打样评估 / 已否决", "候选池", "是", "评审结论", "人工录入"),
]

for i, row in enumerate(fields):
    for c, v in enumerate(row, start=1):
        ws1.cell(row=2 + i, column=c, value=v)

LAST1 = 1 + len(fields)
style_body(ws1, 2, LAST1, 10,
           wrap_cols=(5, 6, 10), center_cols=(1, 3, 4, 7, 9),
           fill_map={3: XL_LAYER_BG, 10: XL_INPUT_BG})

for col, w in zip("ABCDEFGHIJ", [6, 18, 14, 15, 46, 26, 6, 16, 11, 26]):
    ws1.column_dimensions[col].width = w
ws1.freeze_panes = "C2"
ws1.auto_filter.ref = f"A1:J{LAST1}"

# ============ Sheet 2: 评分卡规则 ============
ws2 = wb.create_sheet("评分卡规则")
write_head(ws2, ["序号", "评分维度", "权重", "打分口径", "满分标准", "零分标准", "依据字段"])

rules = [
    (1, "需求缺口", 20, "趋势在涨、但我方产品线没覆盖 = 满分；已有成熟产品覆盖 = 0 分", "趋势品类与我方 SKU 完全不重合", "已有同定位爆品在售", "概念标签、与我方 SKU 重合度、环比增速"),
    (2, "增长动能", 15, "按近 30 天环比增速在该品类中的分位给分", "增速进入品类前 10%", "增速为负", "环比增速、生命周期阶段"),
    (3, "竞争密度", 15, "同价格带 + 同概念的在售商品数越少分越高（逆向指标）", "该价格带该概念无竞品", "已有 20 个以上同质化商品", "价格带、概念标签、榜单排名"),
    (4, "成分与概念壁垒", 15, "是否存在配方工艺、独家原料或专利门槛", "有专利或独家原料", "人人可做的公版配方", "技术壁垒、关键原料可得性"),
    (5, "价格带与客群匹配", 15, "是否落在兰本 / 兰至的主力价格带与目标人群内", "完全落在主力价格带内", "偏离主力价格带一档以上", "价格、与我方价格带匹配、与我方客群匹配"),
    (6, "法规可行性", 10, "备案路径是否顺畅、功效宣称是否撑得住", "普通备案且无需功效评价", "需特殊注册且功效试验周期长", "备案路径、特殊化妆品风险、功效宣称支撑难度"),
    (7, "供应链可行性", 10, "代工能否做、成本与交期是否可接受", "现成产线，成本落在目标毛利率内", "无厂可做或成本远超定价空间", "推荐代工、预估成本带、起订量 MOQ、预估交期（天）"),
]

for i, row in enumerate(rules):
    for c, v in enumerate(row, start=1):
        ws2.cell(row=2 + i, column=c, value=v)

LAST2 = 1 + len(rules)
style_body(ws2, 2, LAST2, 7, wrap_cols=(4, 5, 6, 7), center_cols=(1, 3))

TOTAL2 = LAST2 + 1
ws2.cell(row=TOTAL2, column=2, value="合计")
ws2.cell(row=TOTAL2, column=3, value=100)
ws2.cell(row=TOTAL2, column=4, value="各维度按「权重 × 该维度得分率」加总，得到 0–100 分；调整任一权重后请核对本行合计仍为 100")
for c in range(1, 8):
    cell = ws2.cell(row=TOTAL2, column=c)
    cell.font = F_TOTAL
    cell.fill = PatternFill("solid", fgColor=XL_TOTAL_BG)
    cell.border = BORDER
    cell.alignment = AL_L if c == 4 else AL_C
ws2.cell(row=TOTAL2, column=3).number_format = "0"

for col, w in zip("ABCDEFG", [6, 18, 8, 46, 28, 28, 32]):
    ws2.column_dimensions[col].width = w
ws2.freeze_panes = "C2"
ws2.auto_filter.ref = f"A1:G{LAST2}"

# ============ Sheet 3: 数据源盘点 ============
ws3 = wb.create_sheet("数据源盘点")
write_head(ws3, ["序号", "数据源", "所属区域", "可补层级", "获取方式", "可得性", "费用", "优先级", "备注"])

sources = [
    (1, "抖音电商罗盘", "中国", "③④", "账号导出", "中", "免费", "P0", "需商家账号，品类大盘与竞品销量都能看"),
    (2, "蝉妈妈", "中国", "③④", "网页 + API", "高", "付费", "P0", "单品销量、达人带货数据，接口稳定好接"),
    (3, "飞瓜数据", "中国", "③④", "网页", "高", "付费", "P1", "与蝉妈妈高度重叠，二选一即可"),
    (4, "巨量算数", "中国", "③", "网页", "高", "免费", "P0", "搜索趋势词，判断增速最直接的免费源"),
    (5, "小红书灵犀", "中国", "④", "网页", "中", "免费", "P1", "需品牌号，看口碑与种草密度"),
    (6, "天猫生意参谋", "中国", "③", "账号导出", "中", "免费", "P1", "需店铺账号，看淘系份额"),
    (7, "国家药监局备案平台", "中国", "①②⑤", "网页查询", "高", "免费", "P0", "备案号、成分、企业均可查，权威口径"),
    (8, "美丽修行", "中国", "②④", "网页 + App", "高", "免费 + 会员", "P0", "成分库与安全评分，补②层最快"),
    (9, "KEV 美妆圈（中国区）", "中国", "①②③", "API", "低", "未知", "P0", "你原计划接入的源，接口待打通"),
    (10, "魔镜洞察", "中国", "③", "网页", "高", "付费", "P2", "与蝉妈妈重叠，作为交叉验证"),
    (11, "FastMoss", "跨境", "③④", "网页 + API", "高", "付费", "P0", "TikTok Shop 销量榜，跨境选品主力源"),
    (12, "Kalodata", "跨境", "③", "网页", "高", "付费", "P1", "FastMoss 的交叉验证"),
    (13, "Amazon Best Sellers", "欧美", "③", "网页", "高", "免费", "P0", "榜单免费，看趋势够用"),
    (14, "Keepa", "欧美", "③", "API", "高", "付费", "P1", "需要历史价格与排名曲线时再上"),
    (15, "Sephora / Ulta 榜单", "欧美", "③④", "网页", "中", "免费", "P1", "欧美新品风向标"),
    (16, "Olive Young 榜单", "韩国", "③④", "网页", "中", "免费", "P1", "韩系新品最早上榜的地方"),
    (17, "@cosme 排行", "日本", "③④", "网页", "中", "免费", "P1", "日系口碑基准"),
    (18, "화해 Hwahae", "韩国", "②④", "App", "中", "免费", "P2", "成分与安全评分，与美丽修行同类"),
    (19, "Qoo10 榜单", "日本", "③", "网页", "中", "免费", "P2", "日系电商销量参考"),
    (20, "CosIng（欧盟成分库）", "欧美", "②⑤", "网页", "高", "免费", "P0", "成分法规口径的权威来源"),
    (21, "CPNP / FDA", "欧美", "⑤", "网页", "中", "免费", "P1", "欧美备案路径查询"),
    (22, "MFDS", "韩国", "⑤", "网页", "中", "免费", "P2", "韩系法规对照"),
    (23, "Mintel / WGSN", "全球", "③", "订阅", "高", "付费（高）", "P2", "趋势前瞻，预算充足再上"),
]

for i, row in enumerate(sources):
    for c, v in enumerate(row, start=1):
        ws3.cell(row=2 + i, column=c, value=v)

LAST3 = 1 + len(sources)
style_body(ws3, 2, LAST3, 9, wrap_cols=(9,), center_cols=(1, 3, 4, 5, 6, 7, 8))

for col, w in zip("ABCDEFGHI", [6, 24, 10, 10, 14, 9, 13, 8, 46]):
    ws3.column_dimensions[col].width = w
ws3.freeze_panes = "C2"
ws3.auto_filter.ref = f"A1:I{LAST3}"

# 可得性：低=红，中=黄
ws3.conditional_formatting.add(
    f"F2:F{LAST3}",
    CellIsRule(operator="equal", formula=['"低"'],
               fill=PatternFill("solid", fgColor=XL_RED_BG), font=Font(size=10, color=XL_RED_FG)))
ws3.conditional_formatting.add(
    f"F2:F{LAST3}",
    CellIsRule(operator="equal", formula=['"中"'],
               fill=PatternFill("solid", fgColor=XL_YEL_BG), font=Font(size=10, color=XL_YEL_FG)))

# ============ Sheet 4: 评审说明 ============
ws4 = wb.create_sheet("评审说明")
write_head(ws4, ["类别", "条目", "说明"])

notes = [
    ("基本信息", "用途", "P1 动工前的口径底稿。逐行评审确认后，我按确认版建表、搭页面"),
    ("基本信息", "目标", "给自家品牌（兰本 / 兰至）找新品开发方向，不做铺货型选品"),
    ("基本信息", "市场范围", "全球平铺，先按「数据源盘点」的可得性排优先级"),
    ("基本信息", "模块范围", "P1 三个模块：全球新品库 / 选品评分卡 / 机会看板"),
    ("评审方式", "第 1 步", "看「字段口径表」，不合适的在「评审意见」列写「删」或「改：xxx」"),
    ("评审方式", "第 2 步", "要补的字段直接写「增：字段名 + 口径」，我补进表里"),
    ("评审方式", "第 3 步", "看「数据源盘点」，拿不到的源在备注列标注，我降级为人工录入"),
    ("待定口径", "价格", "海外价按什么汇率折算？含税还是不含税？"),
    ("待定口径", "增速", "环比取 30 天还是 90 天？大促月造成的假峰值要不要剔除？"),
    ("待定口径", "重合度", "「与我方 SKU 重合」以哪条产品线为基准——兰本、兰至，还是合并看？"),
    ("待定口径", "权重", "7 个权重是按「找开发方向」这个目标分配的，改完请核对合计仍为 100"),
    ("评分阈值", "85 分及以上", "进打样评估，直接排期"),
    ("评分阈值", "70 – 84 分", "进候选池，季度复盘时再看"),
    ("评分阈值", "70 分以下", "暂不跟进，留档观察"),
    ("下一步", "确认后", "按确认版在资料库建 5 张云数据表，搭 P1 三个模块，并接入现有情报台"),
]

for i, row in enumerate(notes):
    for c, v in enumerate(row, start=1):
        ws4.cell(row=2 + i, column=c, value=v)

LAST4 = 1 + len(notes)
style_body(ws4, 2, LAST4, 3, wrap_cols=(3,), center_cols=(1, 2))

for col, w in zip("ABC", [12, 16, 78]):
    ws4.column_dimensions[col].width = w
ws4.freeze_panes = "A2"

wb.save("/Users/huizidemac/WorkBuddy/2026-09-19-21-23-47/全球新品选品中台_P1字段口径表.xlsx")
print("OK sheets:", wb.sheetnames)
print("fields:", len(fields), "| rules:", len(rules), "| sources:", len(sources), "| notes:", len(notes))

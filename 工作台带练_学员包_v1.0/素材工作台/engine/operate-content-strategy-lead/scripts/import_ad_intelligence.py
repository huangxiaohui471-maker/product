#!/usr/bin/env python3
from __future__ import annotations

import argparse
import csv
import json
import re
from pathlib import Path
from typing import Any

from common import atomic_write_json, now, sha256_file, stable_id

PROVIDERS = {"feigua_yitou", "appgrowing", "generic"}

ALIASES = {
    "provider_record_id": ["provider_record_id", "素材ID", "广告ID", "创意ID", "记录ID", "ad_id", "record_id"],
    "platform": ["platform", "媒体", "投放平台", "流量平台", "平台"],
    "canonical_url": ["canonical_url", "原视频链接", "抖音链接", "作品链接", "creative_url"],
    "media_url": ["media_url", "下载链接", "视频下载地址", "素材下载地址", "download_url"],
    "source_library_url": ["source_library_url", "素材详情链接", "广告详情链接", "详情链接", "library_url", "source_url"],
    "title": ["title", "素材标题", "广告文案", "文案", "标题", "copy"],
    "advertiser": ["advertiser", "广告主", "品牌", "公司", "商家", "brand"],
    "author_name": ["author_name", "达人", "作者", "账号名称", "抖音号名称"],
    "industry": ["industry", "行业", "推广行业", "品类"],
    "ad_product": ["ad_product", "投放产品", "广告类型", "广告产品", "投放类型"],
    "ad_format": ["ad_format", "素材形式", "广告形式", "创意形式"],
    "objective": ["objective", "推广目标", "投放目标", "转化目标"],
    "first_seen_at": ["first_seen_at", "首次发现", "首次投放", "首次出现"],
    "last_seen_at": ["last_seen_at", "最后发现", "最近投放", "最后出现", "最近出现"],
}

ESTIMATED_ALIASES = {
    "estimated_impressions": ["预估曝光", "预估曝光量", "曝光估算"],
    "exposure_index": ["曝光指数", "曝光热度"],
    "spend_estimate": ["预估消耗", "投放金额估算", "预估投放金额"],
    "placement_count": ["计划数", "关联计划数", "投放计划数"],
    "active_days": ["投放天数", "在投天数", "素材使用天数"],
    "supplier_score": ["跑量分", "投放指数", "素材热度", "增长分"],
}


def _text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _pick(row: dict[str, Any], names: list[str]) -> str:
    for name in names:
        value = _text(row.get(name))
        if value:
            return value
    return ""


def _load(path: Path) -> list[dict[str, Any]]:
    if path.suffix.lower() == ".json":
        payload = json.loads(path.read_text(encoding="utf-8-sig"))
        if isinstance(payload, list):
            return [row for row in payload if isinstance(row, dict)]
        for key in ("rows", "items", "data", "candidates"):
            if isinstance(payload.get(key), list):
                return [row for row in payload[key] if isinstance(row, dict)]
        raise ValueError("JSON 中没有 rows/items/data/candidates 数组")
    if path.suffix.lower() == ".xlsx":
        try:
            from openpyxl import load_workbook
        except ImportError as error:
            raise ValueError("读取 XLSX 需要 openpyxl；也可另存为 UTF-8 CSV") from error
        sheet = load_workbook(path, read_only=True, data_only=True).active
        values = sheet.iter_rows(values_only=True)
        headers = [_text(value) for value in next(values, [])]
        if not any(headers):
            return []
        return [{headers[index]: value for index, value in enumerate(row) if index < len(headers) and headers[index]} for row in values]
    if path.suffix.lower() not in {".csv", ".tsv"}:
        raise ValueError("仅支持 XLSX、UTF-8 CSV/TSV 或 JSON")
    delimiter = "\t" if path.suffix.lower() == ".tsv" else ","
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        return [dict(row) for row in csv.DictReader(handle, delimiter=delimiter)]


def _platform(value: str) -> str:
    text = value.lower()
    if "抖音" in value or "douyin" in text or "千川" in value or "本地推" in value:
        return "douyin"
    # “巨量引擎”同时包含抖音、今日头条、西瓜等版位，不能擅自缩成抖音。
    if "巨量" in value or "oceanengine" in text:
        return "oceanengine_unspecified"
    if "快手" in value or "kuaishou" in text or "磁力" in value:
        return "kuaishou"
    if "小红书" in value or "xiaohongshu" in text or "聚光" in value:
        return "xiaohongshu"
    return text or "unknown"


def _creative_id(provider: str, record_id: str, canonical_url: str, platform: str) -> str:
    if platform == "douyin":
        match = re.search(r"(?:video/|modal_id=)(\d{10,})", canonical_url)
        if match:
            return f"douyin:{match.group(1)}"
    return f"{provider}:{record_id}"


def normalize(row: dict[str, Any], provider: str, row_number: int, verified_provider_export: bool = False) -> dict[str, Any]:
    values = {key: _pick(row, aliases) for key, aliases in ALIASES.items()}
    record_id = values["provider_record_id"] or stable_id(provider, row_number, json.dumps(row, ensure_ascii=False, sort_keys=True))[:20]
    locator = values["canonical_url"] or values["source_library_url"] or values["media_url"]
    if not locator:
        raise ValueError(f"第 {row_number} 行缺少原视频、素材详情或下载链接")
    platform = _platform(values["platform"] or values["ad_product"])
    estimated = {key: _pick(row, aliases) for key, aliases in ESTIMATED_ALIASES.items()}
    estimated = {key: value for key, value in estimated.items() if value != ""}
    return {
        "schema_version": "ad_intelligence_candidate/v1",
        "provider": provider,
        "provider_record_id": record_id,
        "creative_id": _creative_id(provider, record_id, values["canonical_url"], platform),
        "platform": platform,
        "canonical_url": values["canonical_url"],
        "media_url": values["media_url"],
        "source_library_url": values["source_library_url"],
        "title": values["title"],
        "advertiser": values["advertiser"],
        "author_name": values["author_name"],
        "industry": values["industry"],
        "ad_product": values["ad_product"],
        "ad_format": values["ad_format"],
        "objective": values["objective"],
        "first_seen_at": values["first_seen_at"],
        "last_seen_at": values["last_seen_at"],
        # 文件名和 provider 参数不能证明文件真的来自供应商。默认不升格；
        # 只有完成来源验真后由操作者显式确认，才进入第三方广告库证据层。
        "evidence_class": "third_party_ad_library" if verified_provider_export else "user_supplied_unverified",
        "paid_evidence": {
            "level": "strong_signal" if verified_provider_export else "unknown",
            "sources": [f"{provider}_verified_export"] if verified_provider_export else ["unverified_import"],
        },
        "organic_metrics": {},
        "paid_metrics": {},
        "estimated_metrics": estimated,
    }


def import_file(workspace: Path, source: Path, provider: str, verified_provider_export: bool = False) -> Path:
    if provider not in PROVIDERS:
        raise ValueError(f"provider 必须是 {sorted(PROVIDERS)}")
    raw_rows = _load(source)
    candidates, errors = [], []
    for index, row in enumerate(raw_rows, 2):
        try:
            candidates.append(normalize(row, provider, index, verified_provider_export))
        except ValueError as error:
            errors.append({"row": index, "message": str(error)})
    if not candidates:
        raise ValueError(f"没有可导入素材；错误：{errors[:3]}")
    stamp = now().replace(":", "-")
    output = workspace / "runs" / f"adintel-{provider}-{stamp}" / "candidates.json"
    payload = {
        "schema_version": "ad_intelligence_batch/v1",
        "batch_id": stable_id(provider, sha256_file(source))[:24],
        "provider": provider,
        "created_at": now(),
        "source_receipt": {"filename": source.name, "sha256": sha256_file(source), "row_count": len(raw_rows)},
        "evidence_policy": (
            "verified third-party export; third_party_ad_library_max; supplier estimates never become paid_metrics"
            if verified_provider_export else
            "unverified import; user_supplied_unverified; supplier estimates never become paid_metrics"
        ),
        "provider_export_verified": verified_provider_export,
        "candidates": candidates,
        "errors": errors,
        "status": "succeeded" if not errors else "partial",
    }
    atomic_write_json(output, payload)
    return output


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--workspace", required=True)
    parser.add_argument("--source", required=True)
    parser.add_argument("--provider", choices=sorted(PROVIDERS), required=True)
    parser.add_argument(
        "--verified-provider-export",
        action="store_true",
        help="仅在已核验文件确由所选供应商导出时使用；否则保持未验证证据层",
    )
    args = parser.parse_args()
    print(import_file(Path(args.workspace), Path(args.source), args.provider, args.verified_provider_export))
